from io import BytesIO

from docx import Document

from app.text import chunk_text, extract_text, normalize, safe_filename


def test_safe_filename_strips_paths():
    assert safe_filename("../../secrets.pdf") == "secrets.pdf"
    assert safe_filename("") == "document"
    assert safe_filename(None) == "document"


def test_normalize_collapses_blank_lines():
    assert normalize("A\r\n\r\n\r\nB  \n") == "A\n\nB"


def test_short_text_is_one_chunk():
    chunks = chunk_text("Employees may take 18 days of annual leave each calendar year after probation.")
    assert len(chunks) == 1
    assert chunks[0]["index"] == 0


def test_long_text_splits_with_overlap():
    paragraph = "Leave " * 800
    chunks = chunk_text(paragraph, size=500, overlap=80)
    assert len(chunks) > 2
    assert all(len(chunk["text"]) <= 500 for chunk in chunks)
    assert chunks[1]["text"][:40] in paragraph or chunks[0]["text"][-40:] in chunks[1]["text"]


def test_paragraphs_stay_together_when_they_fit():
    text = "Annual leave is 18 days.\n\nSick leave is 10 days.\n\nParental leave is 16 weeks."
    chunks = chunk_text(text, size=200, overlap=20)
    assert len(chunks) == 1
    assert "Sick leave" in chunks[0]["text"]


def test_extract_plain_text():
    assert extract_text("policy.md", b"# Leave\n\nEighteen days.") == "# Leave\n\nEighteen days."


def test_extract_docx_paragraphs_and_table():
    document = Document()
    document.add_paragraph("Code of conduct")
    document.add_paragraph("Treat colleagues with respect.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Receipts required"
    buffer = BytesIO()
    document.save(buffer)
    text = extract_text("conduct.docx", buffer.getvalue())
    assert "Code of conduct" in text
    assert "Receipts required" in text
