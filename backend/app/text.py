import re
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}


def safe_filename(name: str | None) -> str:
    base = Path(name or "document").name.replace("\x00", "").strip()
    if not base or base in {".", ".."}:
        return "document"
    return base[:180]


def extension_of(filename: str) -> str:
    return Path(filename).suffix.lower()


def normalize(text: str) -> str:
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def for_embedding(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def extract_text(filename: str, data: bytes) -> str:
    ext = extension_of(filename)
    if ext == ".pdf":
        return _extract_pdf(data)
    if ext == ".docx":
        return _extract_docx(data)
    if ext in {".txt", ".md", ".markdown"}:
        return _extract_plain(data)
    raise ValueError("Upload a PDF, Word (.docx), Markdown, or text file.")


def _extract_plain(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return normalize(data.decode(encoding))
        except UnicodeDecodeError:
            continue
    return normalize(data.decode("utf-8", errors="replace"))


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(data))
    except Exception as exc:
        raise ValueError("That PDF could not be read. Export it again and retry.") from exc
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ValueError("This PDF is password-protected.") from exc
        if reader.is_encrypted:
            raise ValueError("This PDF is password-protected.")
    parts: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            raw = page.extract_text() or ""
        except Exception:
            raw = ""
        cleaned = raw.strip()
        if cleaned:
            parts.append(f"[Page {index}]\n{cleaned}")
    return normalize("\n\n".join(parts))


def _extract_docx(data: bytes) -> str:
    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise ValueError("That Word file could not be read. Save it as .docx and retry.") from exc
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = []
            seen = set()
            for cell in row.cells:
                value = cell.text.strip()
                if value and value not in seen:
                    seen.add(value)
                    cells.append(value)
            if cells:
                parts.append(" | ".join(cells))
    return normalize("\n\n".join(parts))


def chunk_text(text: str, size: int = 1200, overlap: int = 200) -> list[dict]:
    cleaned = normalize(text)
    if not cleaned:
        return []
    if overlap >= size:
        overlap = max(0, size // 5)

    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", cleaned) if part.strip()]
    chunks: list[str] = []
    current = ""

    def flush() -> None:
        nonlocal current
        piece = current.strip()
        if piece:
            chunks.append(piece)
        current = ""

    for paragraph in paragraphs:
        if len(paragraph) > size:
            flush()
            start = 0
            while start < len(paragraph):
                end = min(start + size, len(paragraph))
                piece = paragraph[start:end].strip()
                if piece:
                    chunks.append(piece)
                if end >= len(paragraph):
                    break
                start = max(end - overlap, start + 1)
            continue

        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if current and len(candidate) > size:
            previous = current
            flush()
            tail = previous[-overlap:].strip() if overlap else ""
            combined = f"{tail}\n\n{paragraph}".strip() if tail else paragraph
            current = combined if len(combined) <= size else paragraph
        else:
            current = candidate

    flush()

    results = []
    for index, piece in enumerate(chunks):
        if len(piece) < 40 and results:
            results[-1]["text"] = f"{results[-1]['text']}\n\n{piece}".strip()
            continue
        results.append({"index": len(results), "text": piece})
    for index, item in enumerate(results):
        item["index"] = index
    return results
